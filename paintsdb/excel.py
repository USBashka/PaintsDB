"""Inventory workbook and lossless monthly price exchange."""
import math
import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .domain import DomainError, month, number, price_values, text


HEADERS = ["Код", "Название", "Период", "Цена упаковки, руб.", "Объём, л", "Масса литра, г/л"]


def cell_code(cell):
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (float, int)) and not isinstance(value, bool):
        if not math.isfinite(value):
            raise DomainError("Некорректный код", "code")
        if float(value).is_integer():
            pattern = cell.number_format
            return str(int(value)).zfill(len(pattern)) if re.fullmatch(r"0+", pattern) else str(int(value))
    return str(value).strip()


def preview(path, override=None):
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise DomainError("Выберите файл XLSX", "path")
    if path.stat().st_size > 25 * 1024 * 1024:
        raise DomainError("Размер файла превышает 25 МБ", "path")
    if override:
        month(override)
    wb = load_workbook(path, data_only=True, read_only=False)
    formulas = load_workbook(path, data_only=False, read_only=False)
    rows, periods, sheets = [], set(), []
    try:
        for ws in wb:
            if ws.max_row > 50000 or ws.max_column > 200:
                raise DomainError("Лист превышает предел 50 000 строк или 200 столбцов")
            header = None
            for row in ws.iter_rows(min_row=1, max_row=min(30, ws.max_row)):
                values = [str(c.value or "").strip() for c in row]
                if "Артикул" in values and "Наименование материала" in values:
                    header = row[0].row
                    mapping = dict(code=values.index("Артикул") + 1, name=values.index("Наименование материала") + 1)
                    for key, prefix in (("volume_liters", "Объём тары"), ("grams_per_liter", "Уд. вес"), ("package_price", "Цена брутто"), ("net", "Вес нетто"), ("unit", "Цена нетто за")):
                        hits = [i + 1 for i, v in enumerate(values) if v.startswith(prefix)]
                        if not hits:
                            raise DomainError(f"Лист {ws.title}: не найден столбец {prefix}")
                        mapping[key] = hits[0]
                    break
                if values[:6] == HEADERS:
                    header = row[0].row
                    mapping = dict(code=1, name=2, period=3, package_price=4, volume_liters=5, grams_per_liter=6)
                    break
            if header is None:
                continue
            sheets.append(ws.title)
            label = " ".join(str(c.value or "") for row in ws.iter_rows(max_row=header) for c in row) + " " + ws.title + " " + path.stem
            match = re.search(r"(0[1-9]|1[0-2])[.](\d{4})", label)
            default_period = override or (f"{match[2]}-{match[1]}" if match else None)
            for idx in range(header + 1, ws.max_row + 1):
                cell = ws.cell(idx, mapping["code"])
                code = cell_code(cell)
                if not code or code in ("Артикул", "Код"):
                    continue
                item = dict(key=f"{ws.title}:{idx}", sheet=ws.title, row=idx, code=code,
                            name=str(ws.cell(idx, mapping["name"]).value or "").strip(), price=None, errors=[], warnings=[])
                try:
                    text(item["code"], "Код", True, 100)
                    text(item["name"], "Название", True, 500)
                    p = override or (str(ws.cell(idx, mapping["period"]).value or "") if "period" in mapping else default_period)
                    month(p)
                    periods.add(p)
                    raw = {k: ws.cell(idx, mapping[k]).value for k in ("package_price", "volume_liters", "grams_per_liter")}
                    if any(v is None or v == "" for v in raw.values()):
                        # Validate present fields even when a monthly price cannot be formed.
                        for k, v in raw.items():
                            if v is not None and v != "":
                                number(v, k, k != "package_price")
                        item["warnings"].append("Нет полного прайса: импортируется только компонент")
                    else:
                        item["price"] = price_values(dict(period=p, **raw))
                        for key, calculated in (("net", "net_grams"), ("unit", "price_per_gram")):
                            if key not in mapping:
                                continue
                            c = ws.cell(idx, mapping[key])
                            if c.value is None:
                                if formulas[ws.title][c.coordinate].data_type == "f":
                                    item["warnings"].append(f"{c.coordinate}: формула без сохранённого результата; использованы исходные поля")
                                continue
                            actual = number(c.value, c.coordinate)
                            decimals = re.search(r"\.([0#]+)", c.number_format)
                            tolerance = 0.5 * 10 ** (-len(decimals[1])) if decimals else 1e-6
                            if abs(actual - item["price"][calculated]) > tolerance + 1e-8:
                                raise DomainError("Производное значение не совпадает с исходными полями", c.coordinate)
                except DomainError as exc:
                    item["errors"].append(f"{exc.field or 'Строка'}: {exc}")
                rows.append(item)
        if not rows:
            raise DomainError("Не найдена таблица компонентов с поддерживаемыми заголовками")
        if not periods:
            raise DomainError("Не удалось однозначно определить месяц. Укажите его перед проверкой", "period")
        counts = Counter(r["code"] for r in rows)
        for r in rows:
            r["duplicate"] = counts[r["code"]] > 1
            if r["duplicate"]:
                r["warnings"].append("Повтор кода: выберите только одну строку")
            r["selected"] = not r["errors"] and not r["duplicate"]
        return dict(period=next(iter(periods)) if len(periods) == 1 else "Несколько месяцев", periods=sorted(periods), rows=rows, sheets=sheets,
                    valid=sum(not r["errors"] for r in rows), errors=sum(bool(r["errors"]) for r in rows))
    finally:
        wb.close(); formulas.close()


def export(path, components, requested_period):
    target = Path(path)
    if target.exists():
        raise DomainError("Файл уже существует. Выберите другое имя", "path")
    wb = Workbook()
    ws = wb.active
    ws.title = "Компоненты"
    ws.append([*HEADERS, "Месяц расчёта"])
    for component in components:
        p = component["price"]
        # Write the actual month so a carried-forward price does not create a fictitious update.
        values = [component["code"], component["name"], p["period"] if p else requested_period,
                  p["package_price"] if p else None, p["volume_liters"] if p else None,
                  p["grams_per_liter"] if p else None, requested_period]
        ws.append(values)
        for c in ws[ws.max_row]:
            if isinstance(c.value, str):
                c.data_type = "s"
        for c in ws[ws.max_row][3:6]:
            c.number_format = "0.000"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for c in ws[1]:
        c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0078D4")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for col, width in zip("ABCDEFG", (20, 65, 15, 24, 18, 24, 18)):
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 32
    with target.open("xb") as stream:
        wb.save(stream)

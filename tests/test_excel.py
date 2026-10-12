from openpyxl import Workbook, load_workbook
from paintsdb.excel import HEADERS
from test_service import service, call, component


def book(path, rows):
    w=Workbook();s=w.active;s.title="2026";s.append(HEADERS)
    for row in rows:s.append(row)
    w.save(path)


def test_import_atomic_repeated_and_aliases(service,tmp_path):
    c=component(service,"007");path=tmp_path/"prices.xlsx"
    book(path,[["007","Новое имя","2026-08",2500,1,1000],["B","Синий","2026-08",5000,2,1000]])
    p=call(service,"preview_import",path=str(path))
    keys=[r["key"] for r in p["rows"]]
    result=call(service,"apply_import",token=p["token"],selected_rows=keys)
    assert result["created"]==1 and result["updated"]==1
    assert call(service,"get_component",id=c["id"])["aliases"]==["алиас"]
    p=call(service,"preview_import",path=str(path))
    result=call(service,"apply_import",token=p["token"],selected_rows=keys)
    assert result["created"]==0 and result["updated"]==2
    assert len(call(service,"list_components"))==2
    assert len(call(service,"get_component",id=c["id"])["prices"])==2


def test_invalid_and_duplicate_rows_require_explicit_selection(service,tmp_path):
    path=tmp_path/"bad.xlsx"
    book(path,[["A","A","2026-08",1000,1,1000],["A","Other","2026-08",2000,1,1000],["B","","2026-08",1000,1,1000],["C","C","2026-08",None,1,None]])
    p=call(service,"preview_import",path=str(path))
    assert not p["rows"][0]["selected"] and not p["rows"][1]["selected"]
    assert p["rows"][2]["errors"] and p["rows"][3]["warnings"]
    result=service.dispatch("apply_import",dict(token=p["token"],selected_rows=[r["key"] for r in p["rows"]]),service.context)
    assert not result["ok"]
    assert call(service,"list_components")==[]
    result=service.dispatch("apply_import",dict(token=p["token"],selected_rows=[p["rows"][0]["key"],p["rows"][1]["key"]]),service.context)
    assert not result["ok"]
    result=call(service,"apply_import",token=p["token"],selected_rows=[p["rows"][0]["key"],p["rows"][3]["key"]])
    assert result["created"]==2 and result["without_price"]==1


def test_export_roundtrip_precision_literal_and_previous_month(service,tmp_path):
    c=component(service,"=DANGEROUS",cost=1234.56789,period="2026-07")
    component(service,"0007",cost=4567.891234,period="2026-08")
    path=tmp_path/"exchange.xlsx"
    call(service,"export_components",path=str(path),period="2026-09")
    w=load_workbook(path); assert w.active['A3'].data_type=='s'
    p=call(service,"preview_import",path=str(path))
    assert p["periods"]==["2026-07","2026-08"]
    assert not any(r["errors"] for r in p["rows"])
    assert {r["code"] for r in p["rows"]}=={"=DANGEROUS","0007"}
    assert next(r for r in p["rows"] if r["code"]=="=DANGEROUS")["price"]["package_price"]==1234.56789


def test_stale_preview(service,tmp_path):
    path=tmp_path/"stale.xlsx";book(path,[["B","B","2026-08",1,1,1]])
    p=call(service,"preview_import",path=str(path));component(service)
    result=service.dispatch("apply_import",dict(token=p["token"],selected_rows=[p["rows"][0]["key"]]),service.context)
    assert result["error"]["code"]=="stale_import"


def test_numeric_code_with_leading_zeros(service,tmp_path):
    path=tmp_path/"codes.xlsx";book(path,[[7,"Семь","2026-08",1,1,1]])
    w=load_workbook(path);w.active['A2'].number_format='0000';w.save(path)
    assert call(service,"preview_import",path=str(path))["rows"][0]["code"]=="0007"
